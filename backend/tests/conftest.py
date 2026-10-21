import io
import pytest
import pymupdf
import docx
from app import create_app
from app.extensions import db as _db
from app.models.user import User
from app.models.job import Job
from app.utils.seed_data import seed_jobs_if_empty
from app.utils.security import generate_auth_token


@pytest.fixture(scope="session")
def app():
    app = create_app("testing")
    return app


@pytest.fixture(scope="function")
def db(app):
    with app.app_context():
        _db.create_all()
        seed_jobs_if_empty(_db, Job)
        yield _db
        _db.session.remove()
        _db.drop_all()


@pytest.fixture
def client(app, db):
    return app.test_client()


@pytest.fixture
def test_user(db):
    user = User(email="testcandidate@example.com", full_name="Alex Turner")
    user.set_password("SecurePassword123!")
    db.session.add(user)
    db.session.commit()
    return user


@pytest.fixture
def auth_headers(test_user):
    token = generate_auth_token(test_user.id)
    return {
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/json"
    }


@pytest.fixture
def second_user(db):
    user = User(email="seconduser@example.com", full_name="Samantha Lee")
    user.set_password("AnotherPassword456!")
    db.session.add(user)
    db.session.commit()
    return user


@pytest.fixture
def second_auth_headers(second_user):
    token = generate_auth_token(second_user.id)
    return {
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/json"
    }


@pytest.fixture
def sample_pdf_bytes():
    """Generates a valid test PDF in memory using PyMuPDF."""
    doc = pymupdf.open()
    page = doc.new_page()
    text = (
        "Alex Turner\n"
        "Email: alex.turner@example.com | Phone: (555) 234-5678\n"
        "LinkedIn: https://linkedin.com/in/alexturner | GitHub: https://github.com/alexturner\n\n"
        "Summary\n"
        "Full stack software engineer with 3 years of experience building web applications in Python, React, and SQL.\n\n"
        "Skills\n"
        "Python, React, JavaScript, TypeScript, Docker, SQL, PostgreSQL, Git, REST API, HTML, CSS\n\n"
        "Experience\n"
        "Software Engineer | Acme Tech (2022 - 2025)\n"
        "• Developed microservices using Flask and PostgreSQL\n"
        "• Built responsive UI interfaces with React and Redux\n\n"
        "Education\n"
        "B.Tech in Computer Science | Stanford University (2018 - 2022)\n\n"
        "Projects\n"
        "Cloud Storage Platform\n"
        "• Built distributed object store with Python and Docker\n"
        "Resume Analyzer App\n"
        "• Designed AI ATS analysis tool with React and FastAPI\n"
    )
    page.insert_text((50, 50), text, fontsize=11)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def sample_docx_bytes():
    """Generates a valid test DOCX in memory using python-docx."""
    doc = docx.Document()
    doc.add_heading("Alex Turner", level=1)
    doc.add_paragraph("Email: alex.turner@example.com | Phone: 555-234-5678")
    doc.add_paragraph("LinkedIn: https://linkedin.com/in/alexturner | GitHub: https://github.com/alexturner")
    doc.add_heading("Technical Skills", level=2)
    doc.add_paragraph("Python, React, SQL, Docker, AWS, Git, Linux, FastApi")
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Full Stack Developer at Nexa Corp (2021 - Present)")
    doc.add_heading("Education", level=2)
    doc.add_paragraph("Bachelor of Science in Computer Science (2017 - 2021)")
    doc.add_heading("Key Projects", level=2)
    doc.add_paragraph("E-Commerce Microservices Engine with Python and Docker")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
