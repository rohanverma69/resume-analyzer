import io
import os
import tempfile
import pymupdf
import docx


class ResumeParseError(Exception):
    """Custom exception raised when resume parsing fails."""
    pass


def extract_text_from_pdf(pdf_bytes: bytes) -> str:
    """
    Extracts text from PDF bytes using PyMuPDF.
    Handles encryption, corrupt streams, and empty text blocks.
    """
    try:
        doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    except Exception as e:
        raise ResumeParseError(f"Unable to read PDF file: {str(e)}")

    try:
        if doc.is_encrypted:
            # Try with empty password
            if not doc.authenticate(""):
                raise ResumeParseError("The PDF document is password-protected or encrypted.")

        if doc.page_count == 0:
            raise ResumeParseError("The PDF document contains 0 pages.")

        extracted_pages = []
        for page_idx in range(doc.page_count):
            page = doc.load_page(page_idx)
            text = page.get_text("text")
            if text:
                extracted_pages.append(text)

        full_text = "\n\n".join(extracted_pages).strip()
        if not full_text:
            raise ResumeParseError(
                "No readable text could be extracted from this PDF. "
                "The document may consist entirely of scanned images without OCR."
            )

        return full_text
    finally:
        doc.close()


def extract_text_from_docx(docx_bytes: bytes) -> str:
    """
    Extracts text from DOCX bytes using python-docx.
    Processes paragraphs, headings, and table cells.
    """
    temp_path = None
    try:
        # Write to secure temporary file
        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as temp_file:
            temp_file.write(docx_bytes)
            temp_path = temp_file.name

        doc = docx.Document(temp_path)
        text_elements = []

        # Paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                text_elements.append(text)

        # Tables
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    # Remove duplicate adjacent cells caused by merged columns
                    deduped = []
                    for t in row_texts:
                        if not deduped or deduped[-1] != t:
                            deduped.append(t)
                    text_elements.append(" | ".join(deduped))

        full_text = "\n".join(text_elements).strip()
        if not full_text:
            raise ResumeParseError(
                "No readable text found in the DOCX document. It may contain only embedded images or empty sections."
            )

        return full_text

    except docx.opc.exceptions.PackageNotFoundError:
        raise ResumeParseError("The uploaded file is not a valid DOCX document or is damaged.")
    except Exception as e:
        if isinstance(e, ResumeParseError):
            raise
        raise ResumeParseError(f"Error parsing DOCX document: {str(e)}")
    finally:
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass


def parse_resume_document(file_bytes: bytes, file_type: str) -> str:
    """
    Main entry point for parsing resume documents.
    Accepts bytes and file_type ('pdf' or 'docx').
    """
    if file_type == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif file_type == "docx":
        return extract_text_from_docx(file_bytes)
    else:
        raise ResumeParseError(f"Unsupported file format: {file_type}")
